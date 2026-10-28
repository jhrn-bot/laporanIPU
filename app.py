import base64
import datetime
import io
import os
import re
import urllib.request
import pandas as pd
import streamlit as st
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_MARKER_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

# ---------------------------------------------------------
# Page Configuration & Constants
# ---------------------------------------------------------
st.set_page_config(
    page_title="Selangor Epi Review Slide Generator",
    page_icon="📊",
    layout="centered",
)

GOOGLE_SLIDES_ID = "1QFVgrEPqgiDditxLQhHRLapQOqaCjZnt"
WABAK_SHEET_ID = "1uVcFp4zSF_gIHdq1BedDFNpuQks0E5AxKJnHbXMsYJE"
BENCANA_SHEET_ID = "1Fp6IORRfdWSJCTC8vqSSoQz6RpCpNXHzO6jj0tHEf2c"
CHART_SHEET_ID = "1SMu8z0MONnxkduZEaRyVNrEnH7KkvnJ9EjuVxSi3WOY"
ILI_SARI_SHEET_ID = "1l9hK5USh1ajbuOC_bwY_R_8T1cmCL-E9c4fqGo5pyTg"

# Clean, validated Base64 QR code image (https://forms.gle/u8M92YMLyQj3rwRT8)
EMBEDDED_QR_B64 = (
    "iVBORw0KGgoAAAANSUhEUgAAATUAAAE1CAYAAACGH3cEAAAPGElEQVR4nO3dMW7cSNPG8eKLPYRhkJfw"
    "Jj2R0z0AOTrFBlS0N3AmJr4EZ+YATh2xo42VDxX4Ak75BvsZ+IIX0yW4tsV5/P8BzAhWs0k9FoFSudm2"
    "bTMAENA0TfOft14EAEQi1ABIIdQASCHUAEgh1ABIIdQASCHUAEgh1ABIIdQASCHUAEgh1ABIIdQASCHU"
    "AEgh1ABIIdQASCHUAEgh1ABI+S3qQtfr1XLOUZe7W8MwFM9ZlsXWdb15Ttu2djgcitc6n8/Fc1JK1nXd"
    "zXO8z89zfx6R74vn/ryi9jOynrrI/TQzsy3I6XTazOyXPzyGYShep+9717U8a5rnOez5RYl8Xzz357XH"
    "eupH8H7y+QlAC6EGQAqhBkAKoQZACqEGQAqhBkAKoQZACqEGQErYXxR4/fHHH/bx48faZX/a58+f7eXl"
    "pVq9y+ViTdMUz/u/fsObjsejPTw83Dyn73vXtTyOx2NYp3zUmvZaz+P9+/f2559/vvUyXu3r16/25cuX"
    "6nWrh9rHjx/tr7/+ql32p53P56qhBvzw7t27u/yZMbM3CTU+PwFIIdQASCHUAEgh1ABIIdQASCHUAEgh"
    "1ABIqd6n5vH09FR1NPg4jq7R2ZH1+r6/eU7O2aZpqrQiv+PxWDxnWZYKK/l3eO6v9vvisSxL1fclpWSP"
    "j4/V6r3GLkMt51x1dnvf91Vf0sPhUKzXNM0uQ019pr7n/mq/Lx7rulZ9Nnv8y4sf+PwEIIVQAyCFUAMg"
    "hVADIIVQAyCFUAMghVADIGWXfWrqlmWxdV2L50RJKRX7irquc/do/epyzsWpxG3b7q6X7VdBqL2BaZqq"
    "Nkp6Or/P57Orm37PTZe1TNNUbIzu+16+UXmv+PwEIIVQAyCFUAMghVADIIVQAyCFUAMghVADIIVQAyCF"
    "UHsDp9PJtm0LOZqmCTk8jbdm5rrW6XT6l3fw9c7ns2vtuH+EGgAphBoAKYQaACmEGgAphBoAKYQaACmE"
    "GgAphBoAKbucfDuOY9Wx0Xscu7wsS3G6qpnZPM/Fc6ZpspzzzXNSSjaOo3t9t+Sc7XK5FOt5JvJ6moLH"
    "cSw+w5SSa68eHh6K5+yR9/6idF1XrdZr7TLUDofDLoOmpnVdXeOgPd37l8ulGGpt27r/qsBTr7R271hw"
    "7/+bUHpfuq5z/SDea6h57+9XwOcnACmEGgAphBoAKYQaACmEGgAphBoAKYQaACnV+9Q+f/7s6j3am+fn"
    "57BrLcti67rePOd6vboakD172bZt8Vpd17muNQxD8Zw9ul6vxV49M3Ptec65+PwiPT8/2++//16tXpRv"
    "3769Sd3qofby8mIvLy+1y+7KNE3FAOn73hUynhHU8zzb09PTzXPO57Or+dbbNLs3Oeew+zsej1X/Yf7+"
    "/bv9/fff1erdOz4/AUgh1ABIIdQASCHUAEgh1ABIIdQASCHUAEgh1ABICWu+HYbhbhsz9+hyubgaa73N"
    "olETXT1rilT7/jw804a9+JmJx29qAKQQagCkEGoApBBqAKQQagCkEGoApBBqAKQQagCkNFtQ99+yLDZN"
    "U/E8T+Pi09OTa/SyR+16Hikla9v25jk5Z9d+esZre8aH19a2rR0Oh+J5nue3LItdr9eb53j3s3YzrGca"
    "7ziOrr2qWc/78+4ReX9N0zS2BTmdTpuZFQ+PYRhc19pjPc8xz3PYfqofUSLfz0hR70vtepHvZ/D98fkJ"
    "QAuhBkAKoQZACqEGQAqhBkAKoQZACqEGQErY5NtIKaWwJsjz+Vw8p9S4afZPs2hKKWJJ1nVdyHW8PM2+"
    "67q6GpD7vo9aVhhPc/GyLFXreZuLPXLOYROHPc/P8362beu61uVyca0rVFTT2z03N3qOvu+rrrt2c+Ne"
    "n5/HHpu1ve9L1Loj7y9S1Pv5inp8fgLQQqgBkEKoAZBCqAGQQqgBkEKoAZBCqAGQQqgB0BLW9eZvjttd"
    "s2Htybf3fETtZ2Qz8z033+5R7f2MZMZvagDEEGoApBBqAKQQagCkEGoApBBqAKQQagCkEGoApDQ/GtZ+"
    "1rIsNk1T8byocdA5Z1c9z+0ty+Ia6R1lmibX6GyPeZ5D6qWUbBzH4rU845k94677vneNWvfwPD/v+zIM"
    "Q/GcyPHop9OpeM7T01PY++KpF/nz4HlfxnEMG33eNE1zt+O8GT/9Nh3wUetmPHrs84t8XyJF7ecr6vH5"
    "CUALoQZACqEGQAqhBkAKoQZACqEGQAqhBkDKb2+9gP/F07y5LEvVepFqNvqa/dMsuhWakL3Nj1HN013X"
    "hTXfppSs67qb57Rt61q7p1k052xN09w8J/L99Dw/r6g9b9vW9c549rz07F4tquntnpsN73mc972q3Qzr"
    "pf78ota913HlZnx+AhBDqAGQQqgBkEKoAZBCqAGQQqgBkEKoAZBCqAGQUj3UmqYpHn3f27ZtNw/PWGKv"
    "0+lUrOc9POZ5Dqvn2c/IvYqqNwxD2H7ukecdfovn51mTZ6T55XJxrb32/ZnxmxoAMYQaACmEGgAphBoA"
    "KYQaACmEGgAphBoAKYQaAClh47xTSjbPc8i1cs7Fscpt27rqHY/H4jnjOBZHEy/LYtM0Fa/lWZN3dHaU"
    "aZqK+5lSssfHx+K1PPfneX7eelG8zy9Kztn17kW9L977i2503aWfmJz7r/GM1/aOEy5dx5zjoCPHlUfy"
    "rMlzRI5nvufnV/uIstdx+lHPz8uMz08AYgg1AFIINQBSCDUAUgg1AFIINQBSCDUAUsKab6/Xq+Wcw661"
    "N23bWt/31epF7mek8/lcPMezV94GZM+ed13nulZUvZyzresaUs+znyml0Hv01NsKU4fXdXW9n7Wfn5nF"
    "df/Vbm6s3bxZG/sZS705tXZz+I6b0fn8BKCFUAMghVADIIVQAyCFUAMghVADIIVQAyCFUAMghVAzs4eH"
    "B2uaJuRQt21b8fCMsfby7HntEdWn06m4B941sZ/xCDUAUgg1AFIINQBSCDUAUgg1AFIINQBSCDUAUgg1"
    "AFKaH9Mif9b1erVlWSIu5eIdJ+wZlzyOo6WUIpZll8vFVa80znqv++nhuT8vTyOoZ7y2dzy658fh6emp"
    "eK22bV3vlKexNrLew8ND8RzP+PDI/YzUNE1Tf95ukMhx15Hjp2vXi7LX/fSIHK8dVc87Hr12vah9itzP"
    "SGZ8fgIQQ6gBkEKoAZBCqAGQQqgBkEKoAZBCqAGQ8lvUhbzNeMMwhNRr29b6vi+e52mGzTmHTa31rGld"
    "V1dTcBRPM6V3Pz1Ktcz874uHZ+2RzcX3zPOMPc3M3mbfNxHV9OZt3qzNs6bIwyOyWdRz0Owb+37ec/Pt"
    "HutFMuPzE4AYQg2AFEINgBRCDYAUQg2AFEINgBRCDYAUQg2AlLBQG4bBtm0rHh7H49Gaprl5RP1lgpnZ"
    "PM+utUfdn0ff91Xrnc/n4p5H/dVFtIeHh+K6PWOzzcy1B55n43mHvUdkvV8Bv6kBkEKoAZBCqAGQQqgB"
    "kEKoAZBCqAGQQqgBkEKoAZASNs7by9MEmVJyjWf2XGue5+I5h8OheM6yLDZNU/E8j8j7q82zpnEci3ua"
    "UnI9myg5Z9fz86wp51wcE3+9Xt1rK5mmKaye5/kty+K61m7FDtMtsx2OZ/bY6/jpe61XW+1x3uoH47wB"
    "oBJCDYAUQg2AFEINgBRCDYAUQg2AFEINgJTqzbelplMzs67riue0beu6VpTIeuu62vl8vnmOtwEyaj8j"
    "5ZyLU1bbtnU1PZf2yeyfRt7SPb7F+5JSqlZvXVfLOVet53k2Hp7n9ypv3Sz3K4ps3owS2XzrObzNm55r"
    "1W729Ty/2s2ptZ9f5BH5/Mz4/AQghlADIIVQAyCFUAMghVADIIVQAyCFUAMghVADICUs1M7nszVN88sf"
    "Ufq+t23biodnTafTqVhvGAZXPY95novX8Xaje9Z0uVyKezAMg6ueZz+jOum99TzPL5Ln+XnX5Hl+0WPr"
    "+U0NgBRCDYAUQg2AFEINgBRCDYAUQg2AFEINgBRCDYCU6uO8379/b+/evatd9qc9Pz/b9+/fq9XLObua"
    "Eud5dl3rcrlELMtVzzOme1kWm6YpYkmWUiqO6l7XNbzJM0LUfuL/iRqj6x0n/OnTp6iSVX348CFsvHbt"
    "cd7q48M946D3Oq48yl73szYzPj8BiCHUAEgh1ABIIdQASCHUAEgh1ABIIdQASKnefOuxLIut61qtXkrJ"
    "uq6rWm8rTJFd19VyzpVWFMvz/JZlCauXcy5OHb5er8UG3dfUK93fuq6hE3JLvPvp2QPP2iOfX7ioprfI"
    "5tvIZlHP4Wk2jGy+9YhsbrznZt+oI7IZdo/3t9fnV5sZn58AxBBqAKQQagCkEGoApBBqAKQQagCkEGoA"
    "pBBqAKQQam/geDxa0zQ3D+/o6dJ1mqaxvu9t27abx+l0+pfv+t8zz3Px/rzd/Z79jPxLgdK6t22zYRjC"
    "6kW9L97DUy/63SPUAEgh1ABIIdQASCHUAEgh1ABIIdQASCHUAEgh1ABI2eU4b/jN81w853A4FM9JKbmu"
    "5TGOY9jobI+cs10ul5vnpJTs8fGxeK2oPagtpWTjOIZcK3I/3wKhdue8f3lQ0nVd2P/TcDgcXEEa5XK5"
    "FLv8t8L/CfFD1H7W1rZt2Noj9/Mt8PkJQAqhBkAKoQZACqEGQAqhBkAKoQZACqEGQMou+9RSSlX7YKL6"
    "s7wi7y9yCmuUlFLVPfXsZ9d1rr3yTJldlsXWdb15Ttu2rl49z5rati02M3v7AqPqeffT04Qd/q5sQU6n"
    "02ZmxePTp09RJav68OGD6/5q86yp9jHPc/V9KPG+nx7DMBSv0/e961q19zOqXuR+RjLj8xOAGEINgBRC"
    "DYAUQg2AFEINgBRCDYAUQg2AlOrNt1+/fq1dMsS3b9/eegkAHKqH2pcvX+zLly+1y+7K8XgM+0uAzfGX"
    "CbXrRWqapnjOPM/Fqa/DMLjW7qkXqfZ+1q73Fvj8BCCFUAMghVADIIVQAyCFUAMghVADIIVQAyCFUAMg"
    "Jaz5NqVk8zxHXU7aOI6uMcf3Wi+S553yjrKOqudRe0R8bXv+eW+2X6HFGMAvoWmahs9PAFIINQBSCDUA"
    "Ugg1AFIINQBSCDUAUgg1AFIINQBSCDUAUgg1AFIINQBSCDUAUgg1AFIINQBSCDUAUgg1AFIINQAAgL36"
    "LylpEZX5Z/bkAAAAAElFTkSuQmCC"
)

VALID_DISTRICTS = [
    "GOMBAK",
    "HULU LANGAT",
    "HULU SELANGOR",
    "KLANG",
    "KUALA LANGAT",
    "KUALA SELANGOR",
    "PETALING",
    "SABAK BERNAM",
    "SEPANG",
]

DISTRICT_ABBR = [
    ("GOMBAK", "GBK"),
    ("HULU LANGAT", "HL"),
    ("HULU SELANGOR", "HS"),
    ("KLANG", "KLG"),
    ("KUALA LANGAT", "KL"),
    ("KUALA SELANGOR", "KS"),
    ("PETALING", "PTG"),
    ("SABAK BERNAM", "SB"),
    ("SEPANG", "SPG"),
]

INCLUSION_DIAGNOSES = (
    "HFMD",
    "FOOD POISONING",
    "COVID-19",
    "MERS-COV",
    "DIPHTERIA",
    "DIPHTHERIA",
    "CHIKUNGUNYA",
    "MALARIA",
    "LEPROSY",
)

EXCLUSION_DIAGNOSES_14 = (
    "HFMD",
    "FOOD POISONING",
    "COVID-19",
    "MERS-COV",
    "DIPHTERIA",
    "DIPHTHERIA",
    "CHIKUNGUNYA",
    "MALARIA",
    "LEPROSY",
    "DENGUE/DHF",
    "DENGUE",
    "MEASLES",
    "HIV/AIDS",
    "HIV",
)

DIAG_TEMPOH_24H = [
    "DENGUE/DHF",
    "DENGUE",
    "COVID-19",
    "HFMD",
    "MEASLES",
    "AVIAN INFLUENZA",
    "FOOD POISONING",
    "MERS-COV",
    "MALARIA",
    "DIPHTERIA",
    "DIPHTHERIA",
    "CHIKUNGUNYA",
    "RABIES",
    "MONKEYPOX",
    "MPOX",
    "POLIOMYELITIS",
    "EBOLA",
    "CHOLERA",
    "ZIKA VIRUS INFECTION",
    "YELLOW FEVER",
    "PLAGUE",
]

DIAG_TEMPOH_7D = [
    "GONORRHOEA",
    "HIV/AIDS",
    "HIV",
    "VIRAL HEPATITIS",
    "PERTUSSIS",
    "LEPTOSPIROSIS",
    "SYPHILIS",
    "TUBERCULOSIS",
    "RELAPSING FEVER",
    "TYPHOID/PARATYPHOID",
    "DYSENTRY",
    "LEPROSY",
    "VIRAL ENCEPHALITIS",
    "TETANUS",
    "TYPHUS",
    "CHANCROID",
    "WHOOPING COUGH",
    "BRUCELLOSIS",
    "MELIODOSIS",
    "OTHER SPECIFIED VIRAL HEPATITIS",
]

NAVY = RGBColor(27, 54, 93)
LIGHT_GREY = RGBColor(211, 211, 211)

DISEASE_COLORS = {
    "Denggi": RGBColor(255, 0, 0),
    "Malaria": RGBColor(237, 125, 49),
    "Chikungunya": RGBColor(146, 208, 80),
    "HFMD": RGBColor(169, 209, 142),
    "Keracunan Makanan": RGBColor(68, 114, 196),
    "Influenza / ILI": RGBColor(96, 40, 130),
    "Chickenpox": RGBColor(0, 32, 96),
    "COVID-19": RGBColor(255, 0, 0),
    "Tuberkulosis": RGBColor(112, 48, 160),
    "Measles": RGBColor(128, 96, 0),
    "Rotavirus": RGBColor(55, 86, 35),
    "Norovirus": RGBColor(122, 182, 229),
    "Disyaki Norovirus": RGBColor(248, 161, 108),
    "AGE": RGBColor(180, 180, 180),
    "Respiratory Syncytial Virus (RSV)": RGBColor(139, 58, 0),
    "Pertussis": RGBColor(200, 200, 200),
    "Scabies": RGBColor(84, 130, 53),
    "Hepatitis A": RGBColor(146, 208, 80),
    "Adenovirus": RGBColor(47, 117, 181),
    "Difteria": RGBColor(220, 88, 21),
    "Leptospirosis": RGBColor(127, 127, 127),
    "Disentri": RGBColor(197, 143, 0),
    "Disyaki Rotavirus": RGBColor(46, 85, 155),
    "Mpox": RGBColor(84, 130, 53),
    "Konjunktivitis": RGBColor(157, 195, 230),
}

SAMPEL_COLORS = {
    "INFLUENZA A": RGBColor(146, 208, 80),
    "INFLUENZA B": RGBColor(68, 114, 196),
    "COVID-19": RGBColor(255, 192, 0),
}

HEADER_CLEAN_MAP = {
    "denggi": "Denggi",
    "malaria": "Malaria",
    "hikungun": "Chikungunya",
    "chikungun": "Chikungunya",
    "hfmd": "HFMD",
    "unan mai": "Keracunan Makanan",
    "keracunan": "Keracunan Makanan",
    "luenza": "Influenza / ILI",
    "influenza": "Influenza / ILI",
    "hickenpo": "Chickenpox",
    "chickenpox": "Chickenpox",
    "covid": "COVID-19",
    "berkulos": "Tuberkulosis",
    "tuberkulosis": "Tuberkulosis",
    "measles": "Measles",
    "rotavirus": "Rotavirus",
    "orovirus": "Norovirus",
    "norovirus": "Norovirus",
    "aki noro": "Disyaki Norovirus",
    "disyaki noro": "Disyaki Norovirus",
    "age": "AGE",
    "syncytia": "Respiratory Syncytial Virus (RSV)",
    "rsv": "Respiratory Syncytial Virus (RSV)",
    "pertussis": "Pertussis",
    "scabies": "Scabies",
    "epatitis": "Hepatitis A",
    "hepatitis": "Hepatitis A",
    "denoviru": "Adenovirus",
    "adenovirus": "Adenovirus",
    "difteria": "Difteria",
    "diphtheria": "Difteria",
    "ptospiro": "Leptospirosis",
    "leptospirosis": "Leptospirosis",
    "disentri": "Disentri",
    "aki rota": "Disyaki Rotavirus",
    "disyaki rota": "Disyaki Rotavirus",
    "mpox": "Mpox",
    "junktivitis": "Konjunktivitis",
    "konjunktivitis": "Konjunktivitis",
}


# ---------------------------------------------------------
# Helper Functions
# ---------------------------------------------------------
def get_previous_epi_week():
  myt_zone = datetime.timezone(datetime.timedelta(hours=8))
  today = datetime.datetime.now(myt_zone).date()
  last_week_date = today - datetime.timedelta(days=7)
  jan_1 = datetime.date(last_week_date.year, 1, 1)
  jan_1_day = jan_1.isoweekday() % 7
  first_sunday = (
      jan_1
      if jan_1_day == 0
      else jan_1 + datetime.timedelta(days=(7 - jan_1_day))
  )
  if last_week_date < first_sunday:
    return 1, last_week_date.year
  days_diff = (last_week_date - first_sunday).days
  return (days_diff // 7) + 1, last_week_date.year


epi_week, year = get_previous_epi_week()


def parse_excel_date(val):
  if pd.isna(val) or str(val).strip().upper() in ["NAN", ""]:
    return None
  if isinstance(val, (datetime.datetime, datetime.date, pd.Timestamp)):
    return val.date() if hasattr(val, "date") else val
  try:
    return (
        datetime.datetime(1899, 12, 30) + datetime.timedelta(days=float(val))
    ).date()
  except:
    pass
  try:
    return pd.to_datetime(val, errors="coerce").date()
  except:
    return None


def set_cell_border(cell, color=NAVY, width=Pt(1.5)):
  tcPr = cell._tc.get_or_add_tcPr()
  for line_type in ["a:lnL", "a:lnR", "a:lnT", "a:lnB"]:
    ln = tcPr.find(qn(line_type))
    if ln is None:
      ln = OxmlElement(line_type)
      tcPr.append(ln)
    ln.set("w", str(int(width)))
    solidFill = OxmlElement("a:solidFill")
    srgbClr = OxmlElement("a:srgbClr")
    srgbClr.set("val", f"{color[0]:02X}{color[1]:02X}{color[2]:02X}")
    solidFill.append(srgbClr)
    ln.append(solidFill)


def write_cell(cell, text, bold=True, align_left=False, is_red=False, size=12):
  cell.text = ""
  tf = cell.text_frame
  tf.word_wrap = True
  for idx, line_str in enumerate(str(text).split("\n")):
    p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
    p.alignment = PP_ALIGN.LEFT if align_left else PP_ALIGN.CENTER
    run = p.add_run()
    run.text = line_str
    run.font.size = Pt(size)
    run.font.name = "Calibri"
    run.font.bold = bold
    run.font.color.rgb = RGBColor(255, 0, 0) if is_red else NAVY
  return tf.paragraphs[0]


def write_wabak_cell(cell, text, bold=True, align_left=False, size=13):
  cell.text = ""
  tf = cell.text_frame
  tf.word_wrap = True
  p = tf.paragraphs[0]
  p.alignment = PP_ALIGN.LEFT if align_left else PP_ALIGN.CENTER
  val_str = str(text).strip()
  if val_str.lower() in ["nan", "none", ""]:
    val_str = "-"
  val_str = re.sub(r"\(Rsv\)", "(RSV)", val_str, flags=re.IGNORECASE)
  match = re.search(r"^(.*?)\s*(\([0-9]+\))$", val_str)
  if match:
    if match.group(1).strip():
      run1 = p.add_run()
      run1.text = match.group(1).strip() + " "
      (
          run1.font.size,
          run1.font.name,
          run1.font.bold,
          run1.font.color.rgb,
      ) = (
          Pt(size),
          "Calibri",
          bold,
          NAVY,
      )
    run2 = p.add_run()
    run2.text = match.group(2).strip()
    (
        run2.font.size,
        run2.font.name,
        run2.font.bold,
        run2.font.color.rgb,
    ) = (
        Pt(size),
        "Calibri",
        bold,
        RGBColor(255, 0, 0),
    )
  else:
    run = p.add_run()
    run.text = val_str
    (
        run.font.size,
        run.font.name,
        run.font.bold,
        run.font.color.rgb,
    ) = (
        Pt(size),
        "Calibri",
        bold,
        NAVY,
    )


def write_header_with_red_me(cell, prefix, me_str, size=13, newline=False):
  cell.text = ""
  p = cell.text_frame.paragraphs[0]
  p.alignment = PP_ALIGN.CENTER
  run1 = p.add_run()
  run1.text = prefix + ("\n" if newline else " ")
  (
      run1.font.size,
      run1.font.name,
      run1.font.bold,
      run1.font.color.rgb,
  ) = (
      Pt(size),
      "Calibri",
      True,
      NAVY,
  )
  run2 = p.add_run()
  run2.text = me_str
  (
      run2.font.size,
      run2.font.name,
      run2.font.bold,
      run2.font.color.rgb,
  ) = (
      Pt(size),
      "Calibri",
      True,
      RGBColor(255, 0, 0),
  )


format_cell_stat = (
    lambda tot, swasta, zero_as_dash=True: (
        "-"
        if tot == 0 and zero_as_dash
        else (
            "0"
            if tot == 0
            else f"{tot} ({swasta})" if swasta > 0 else f"{tot}"
        )
    )
)


def remove_axis_line(axis_elem):
  spPr = axis_elem.find(qn("c:spPr"))
  if spPr is None:
    spPr = OxmlElement("c:spPr")
    axis_elem.append(spPr)
  ln = spPr.find(qn("a:ln"))
  if ln is None:
    ln = OxmlElement("a:ln")
    spPr.append(ln)
  noFill = ln.find(qn("a:noFill"))
  if noFill is None:
    ln.append(OxmlElement("a:noFill"))


def set_val_axis_title_oxml(
    chart, title_text, axis_index=0, font_size=10, bold=True
):
  plotArea = chart.element.find(qn("c:chart")).find(qn("c:plotArea"))
  valAxes = plotArea.findall(qn("c:valAx"))
  if len(valAxes) <= axis_index:
    return
  valAx = valAxes[axis_index]

  title_elem = valAx.find(qn("c:title"))
  if title_elem is None:
    title_elem = OxmlElement("c:title")
    valAx.append(title_elem)
  else:
    title_elem.clear()

  tx = OxmlElement("c:tx")
  rich = OxmlElement("c:rich")
  bodyPr = OxmlElement("a:bodyPr")
  bodyPr.set("rot", "-5400000")
  bodyPr.set("vert", "horz")
  rich.append(bodyPr)
  rich.append(OxmlElement("a:lstStyle"))

  p = OxmlElement("a:p")
  pPr = OxmlElement("a:pPr")
  defRPr = OxmlElement("a:defRPr")
  defRPr.set("sz", str(int(font_size * 100)))
  defRPr.set("b", "1" if bold else "0")
  defRPr.set("bld", "1" if bold else "0")
  pPr.append(defRPr)
  p.append(pPr)

  r = OxmlElement("a:r")
  rPr = OxmlElement("a:rPr")
  rPr.set("sz", str(int(font_size * 100)))
  rPr.set("b", "1" if bold else "0")
  rPr.set("bld", "1" if bold else "0")

  solidFill = OxmlElement("a:solidFill")
  srgbClr = OxmlElement("a:srgbClr")
  srgbClr.set("val", "1B365D")
  solidFill.append(srgbClr)
  rPr.append(solidFill)

  r.append(rPr)
  t = OxmlElement("a:t")
  t.text = title_text
  r.append(t)
  p.append(r)
  rich.append(p)
  tx.append(rich)
  title_elem.append(tx)


def process_sampel_dataframe(df_sampel, target_epi_week, target_year):
  if df_sampel.empty:
    return pd.DataFrame(), -1

  col0 = df_sampel.columns[0]
  parsed_rows = []
  current_year = target_year - 1

  for idx, row in df_sampel.iterrows():
    val = str(row[col0]).strip()
    w_num = None
    y_num = None

    if "/" in val:
      parts = val.split("/")
      if len(parts) == 2:
        try:
          w_num = int(parts[0])
          y_num = int(parts[1])
          if y_num < 100:
            y_num += 2000
        except ValueError:
          pass
      elif len(parts) == 3:
        try:
          d = pd.to_datetime(val, dayfirst=True, errors="coerce")
          if not pd.isna(d):
            w_num = int(d.isocalendar().week)
            y_num = int(d.year)
        except Exception:
          pass
    else:
      try:
        w_num = int(float(val))
      except ValueError:
        pass

    if y_num is None and w_num is not None:
      if idx > 0 and len(parsed_rows) > 0:
        prev_w = parsed_rows[-1]["w_num"]
        prev_y = parsed_rows[-1]["y_num"]
        if prev_w is not None and prev_y is not None:
          if w_num < prev_w and prev_w >= 40:
            current_year = prev_y + 1
          else:
            current_year = prev_y
      y_num = current_year
    elif y_num is not None:
      current_year = y_num

    if w_num is not None and y_num is not None:
      parsed_rows.append({
          "idx": idx,
          "w_num": w_num,
          "y_num": y_num,
          "label": f"{w_num}/{y_num}",
          "row": row,
      })

  keep_rows = []
  idx_reset = -1

  for i, p in enumerate(parsed_rows):
    if p["y_num"] == target_year and p["w_num"] > target_epi_week:
      break

    if i > 0 and p["y_num"] > parsed_rows[i - 1]["y_num"]:
      idx_reset = i

    keep_rows.append(p)

  if not keep_rows:
    return pd.DataFrame(), -1

  clean_df = pd.DataFrame([k["row"] for k in keep_rows])
  clean_df[col0] = [k["label"] for k in keep_rows]

  return clean_df.reset_index(drop=True), idx_reset


def format_prestasi_val(val, col_idx, row_idx):
  if pd.isna(val) or str(val).strip().upper() in ["NAN", "NONE", ""]:
    return "-" if row_idx > 0 and col_idx >= 2 else ""

  val_str = str(val).strip()

  if col_idx in [3, 4] and row_idx > 0:
    if val_str.endswith("%"):
      return val_str
    try:
      num = float(val_str)
      if num <= 1.0:
        return f"{int(round(num * 100))}%"
      else:
        return f"{int(round(num))}%"
    except ValueError:
      return val_str

  if col_idx == 2 and row_idx > 0:
    try:
      return f"{int(round(float(val_str))):,}"
    except ValueError:
      return val_str

  return val_str


# ---------------------------------------------------------
# Data Fetchers
# ---------------------------------------------------------
def fetch_google_slides_pptx(file_id):
  req = urllib.request.Request(
      f"https://docs.google.com/presentation/d/{file_id}/export/pptx",
      headers={"User-Agent": "Mozilla/5.0"},
  )
  with urllib.request.urlopen(req) as response:
    return io.BytesIO(response.read())


@st.cache_data(ttl=600)
def fetch_wabak_data():
  req = urllib.request.Request(
      f"https://docs.google.com/spreadsheets/d/{WABAK_SHEET_ID}/gviz/tq?tqx=out:csv&sheet=final",
      headers={"User-Agent": "Mozilla/5.0"},
  )
  try:
    with urllib.request.urlopen(req) as resp:
      df_sheet = pd.read_csv(io.BytesIO(resp.read()), header=None)
      df_p_aa = df_sheet.iloc[1:, 15:27].copy()
      df_p_aa = df_p_aa.dropna(subset=[df_p_aa.columns[1]])
      return df_p_aa[df_p_aa.iloc[:, 1].astype(str).str.strip() != ""]
  except:
    return pd.DataFrame()


@st.cache_data(ttl=600)
def fetch_bencana_data():
  req = urllib.request.Request(
      f"https://docs.google.com/spreadsheets/d/{BENCANA_SHEET_ID}/gviz/tq?tqx=out:csv&sheet=table%202026",
      headers={"User-Agent": "Mozilla/5.0"},
  )
  try:
    with urllib.request.urlopen(req) as resp:
      df_sheet = pd.read_csv(io.BytesIO(resp.read()), header=None)
      df_ah_au = df_sheet.iloc[2:, 33:47].copy()
      df_ah_au = df_ah_au.dropna(subset=[df_ah_au.columns[0]])
      return df_ah_au[df_ah_au.iloc[:, 0].astype(str).str.strip() != ""]
  except:
    return pd.DataFrame()


def fetch_graf_data():
  url = (
      f"https://docs.google.com/spreadsheets/d/{CHART_SHEET_ID}/gviz/tq?tqx=out:csv&sheet=GRAF%20WABAK%20S2WER&range=B2:AD120"
  )
  req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
  try:
    with urllib.request.urlopen(req) as resp:
      raw_df = pd.read_csv(io.BytesIO(resp.read()), header=None)

      if raw_df.empty or len(raw_df) < 2:
        return pd.DataFrame()

      raw_headers = [str(x).strip() for x in raw_df.iloc[0].tolist()]
      data_df = raw_df.iloc[1:].copy()

      data_df[0] = pd.to_numeric(data_df[0], errors="coerce")
      data_df = data_df.dropna(subset=[0])
      data_df[0] = data_df[0].astype(int).astype(str)

      clean_dict = {"Minggu Epid": data_df[0].tolist()}

      for col_i in range(1, len(raw_headers)):
        raw_name = raw_headers[col_i]
        clean_name = raw_name
        raw_name_low = str(raw_name).lower().strip()

        if (
            raw_name_low in ["nan", "none", "", "null"]
            or "unnamed" in raw_name_low
        ):
          continue

        for key, val in HEADER_CLEAN_MAP.items():
          if key in raw_name_low:
            clean_name = val
            break
        vals = (
            pd.to_numeric(data_df[col_i], errors="coerce").fillna(0).tolist()
        )
        clean_dict[clean_name] = vals

      return pd.DataFrame(clean_dict)
  except Exception as e:
    st.warning(f"Note: Could not retrieve chart data: {e}")
    return pd.DataFrame()


def fetch_ili_sari_data():
  url = (
      f"https://docs.google.com/spreadsheets/d/{ILI_SARI_SHEET_ID}/gviz/tq?tqx=out:csv&sheet=Sheet23&range=AM1:AP120"
  )
  req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
  try:
    with urllib.request.urlopen(req) as resp:
      raw_df = pd.read_csv(io.BytesIO(resp.read()), header=None)

      if raw_df.empty or len(raw_df) < 2:
        return pd.DataFrame()

      raw_headers = [str(x).strip() for x in raw_df.iloc[0].tolist()]
      data_df = raw_df.iloc[1:].copy()

      data_df[0] = pd.to_numeric(data_df[0], errors="coerce")
      data_df = data_df.dropna(subset=[0])
      data_df[0] = data_df[0].astype(int).astype(str)

      clean_dict = {"Minggu Epid": data_df[0].tolist()}

      for col_i in range(1, len(raw_headers)):
        h_name = (
            raw_headers[col_i] if col_i < len(raw_headers) else f"Col_{col_i}"
        )
        if (
            str(h_name).lower().strip() in ["nan", "none", "", "null"]
            or "unnamed" in str(h_name).lower()
        ):
          continue
        vals = pd.to_numeric(data_df[col_i], errors="coerce").fillna(0).tolist()
        clean_dict[h_name] = vals

      return pd.DataFrame(clean_dict)
  except Exception as e:
    st.warning(f"Note: Could not retrieve ILI/SARI chart data: {e}")
    return pd.DataFrame()


def fetch_survelan_sampel_data():
  urls = [
      (
          f"https://docs.google.com/spreadsheets/d/{ILI_SARI_SHEET_ID}/gviz/tq?tqx=out:csv&sheet=Sheet23&range=AR1:AU120"
      ),
      (
          f"https://docs.google.com/spreadsheets/d/{ILI_SARI_SHEET_ID}/gviz/tq?tqx=out:csv&sheet=ILIMakmal&range=C1:F120"
      ),
  ]
  for url in urls:
    try:
      req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
      with urllib.request.urlopen(req) as resp:
        raw_df = pd.read_csv(io.BytesIO(resp.read()), header=None)

        if raw_df.empty or len(raw_df) < 2:
          continue

        raw_headers = [str(x).strip() for x in raw_df.iloc[0].tolist()]
        data_df = raw_df.iloc[1:].copy()

        cat_vals = [str(x).strip() for x in data_df[0].tolist()]
        clean_dict = {"ME/TAHUN": cat_vals}

        default_series = ["INFLUENZA A", "INFLUENZA B", "COVID-19"]

        for col_i in range(1, len(data_df.columns)):
          raw_h = (
              raw_headers[col_i].strip()
              if col_i < len(raw_headers)
              else f"Col_{col_i}"
          )
          raw_upper = raw_h.upper()

          if "INFLUENZA A" in raw_upper or "INF A" in raw_upper:
            col_name = "INFLUENZA A"
          elif "INFLUENZA B" in raw_upper or "INF B" in raw_upper:
            col_name = "INFLUENZA B"
          elif "COVID" in raw_upper:
            col_name = "COVID-19"
          else:
            col_name = (
                default_series[col_i - 1]
                if (col_i - 1) < len(default_series)
                else f"Series_{col_i}"
            )

          vals = (
              pd.to_numeric(data_df[col_i], errors="coerce").fillna(0).tolist()
          )
          clean_dict[col_name] = vals

        res_df = pd.DataFrame(clean_dict)
        res_df = res_df[
            res_df["ME/TAHUN"]
            .astype(str)
            .str.strip()
            .str.upper()
            .isin(["NAN", "NONE", "", "NULL"])
            == False
        ]
        if not res_df.empty:
          return res_df.reset_index(drop=True)
    except Exception:
      continue
  return pd.DataFrame()


def fetch_prestasi_klinik_data():
  url = f"https://docs.google.com/spreadsheets/d/{ILI_SARI_SHEET_ID}/gviz/tq?tqx=out:csv&sheet=S2WER&range=M35:Q42"
  req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
  try:
    with urllib.request.urlopen(req) as resp:
      raw_df = pd.read_csv(io.BytesIO(resp.read()), header=None)
      if not raw_df.empty:
        return raw_df
  except Exception as e:
    st.warning(f"Note: Could not retrieve Prestasi Klinik data: {e}")
  return pd.DataFrame()


# ---------------------------------------------------------
# Main Data Processing
# ---------------------------------------------------------
@st.cache_data
def load_and_process_data(file, target_me, inclusion_tuple, exclusion_tuple):
  df = pd.read_excel(file)
  (c_me, c_bc, c_bi, c_bp, c_br, c_bt, c_cd, c_dt, c_dx, c_dy) = (
      df.columns[5],
      df.columns[54],
      df.columns[60],
      df.columns[67],
      df.columns[69],
      df.columns[71],
      df.columns[81],
      df.columns[123],
      df.columns[127],
      df.columns[128],
  )

  df_clean = df[df[c_dt].isin(VALID_DISTRICTS)].copy()
  df_clean[c_me] = pd.to_numeric(df_clean[c_me], errors="coerce")
  df_semasa = df_clean[df_clean[c_me] == target_me]
  df_kumulatif = df_clean[df_clean[c_me] <= target_me]

  def get_stats(d_sub):
    tot = len(d_sub)
    c = d_sub[c_br].value_counts().to_dict()
    return {
        "total": tot,
        "daftar_notifikasi": c.get("Daftar Notifikasi", 0),
        "daftar_kes": c.get("Daftar Kes", 0),
        "abai": c.get("Abai Notifikasi", 0),
        "belum": c.get("Belum Ambil Tindakan", 0),
        "batal": c.get("Batal Daftar", 0),
        "pct_daftar_notif": (
            f"{(c.get('Daftar Notifikasi', 0)/tot*100):.0f}%" if tot else "0%"
        ),
        "pct_daftar_kes": (
            f"{(c.get('Daftar Kes', 0)/tot*100):.0f}%" if tot else "0%"
        ),
        "pct_abai": (
            f"{(c.get('Abai Notifikasi', 0)/tot*100):.0f}%" if tot else "0%"
        ),
        "pct_belum": (
            f"{(c.get('Belum Ambil Tindakan', 0)/tot*100):.2f}%"
            if tot
            else "0.00%"
        ),
        "pct_batal": (
            f"{(c.get('Batal Daftar', 0)/tot*100):.2f}%" if tot else "0.00%"
        ),
    }

  stats_semasa = get_stats(df_semasa)
  stats_kumulatif = get_stats(df_kumulatif)

  df_clean[c_dx] = (
      df_clean[c_dx]
      .astype(str)
      .str.strip()
      .str.upper()
      .replace({"MONKEYPOX": "MPOX"})
  )
  df_kumu_penyakit = df_clean[df_clean[c_me] <= target_me]

  data_penyakit = []
  for dx in df_kumu_penyakit[c_dx].unique():
    if pd.isna(dx) or dx == "NAN":
      continue
    d_dx = df_kumu_penyakit[df_kumu_penyakit[c_dx] == dx]
    tot_notif = len(d_dx)
    if tot_notif == 0:
      continue
    is_dk = d_dx[c_br].astype(str).str.strip().str.upper() == "DAFTAR KES"
    is_m = d_dx[c_bt].astype(str).str.strip().str.upper() == "MATI"
    data_penyakit.append({
        "Penyakit": dx,
        "Semasa": len(d_dx[(d_dx[c_me] == target_me) & is_dk]),
        "Kumulatif": len(d_dx[is_dk]),
        "Mati": len(d_dx[is_m]),
        "Peratus": len(d_dx[is_dk]) / tot_notif * 100,
    })
  df_penyakit = (
      pd.DataFrame(data_penyakit).sort_values(
          by=["Kumulatif", "Penyakit"], ascending=[False, True]
      )
      if data_penyakit
      else pd.DataFrame()
  )

  df_district_data = []
  for d in VALID_DISTRICTS:
    d_d = df_kumulatif[df_kumulatif[c_dt] == d]
    jml = len(d_d)
    c = d_d[c_br].value_counts().to_dict()
    df_district_data.append({
        "DAERAH": d,
        "Jumlah Notifikasi": jml,
        "Daftar Notifikasi": c.get("Daftar Notifikasi", 0),
        "pct_notif": c.get("Daftar Notifikasi", 0) / jml * 100 if jml else 0,
        "Daftar Kes": c.get("Daftar Kes", 0),
        "pct_kes": c.get("Daftar Kes", 0) / jml * 100 if jml else 0,
        "Abai Notifikasi": c.get("Abai Notifikasi", 0),
        "pct_abai": c.get("Abai Notifikasi", 0) / jml * 100 if jml else 0,
        "Batal Daftar": c.get("Batal Daftar", 0),
        "pct_batal": c.get("Batal Daftar", 0) / jml * 100 if jml else 0,
        "Belum Ambil Tindakan": c.get("Belum Ambil Tindakan", 0),
        "pct_belum": c.get("Belum Ambil Tindakan", 0) / jml * 100 if jml else 0,
    })
  df_district = (
      pd.DataFrame(df_district_data).sort_values(
          by="Jumlah Notifikasi", ascending=False
      )
      if df_district_data
      else pd.DataFrame()
  )

  def make_ct(cond):
    d_sub = df_clean[cond].copy()
    if d_sub.empty:
      return pd.DataFrame()
    ct = pd.crosstab(
        d_sub[c_dx] if c_dx in d_sub else d_sub[c_dy], d_sub[c_dt]
    ).reindex(columns=[x[0] for x in DISTRICT_ABBR], fill_value=0)
    ct["JUM"] = ct.sum(axis=1)
    return (
        ct[ct["JUM"] > 0]
        .sort_values(
            by=["JUM", ct.index.name or "index"], ascending=[False, True]
        )
        .reset_index()
    )

  cond_belum = (df_clean[c_me] == target_me) & (
      df_clean[c_br].astype(str).str.strip().str.upper()
      == "BELUM AMBIL TINDAKAN"
  )
  df_belum_ct = make_ct(cond_belum)

  df_clean[c_dy] = (
      df_clean[c_dy]
      .fillna("")
      .astype(str)
      .str.strip()
      .apply(lambda x: "Tiada subdiagnosis" if x in ["", "NAN"] else x)
  )
  df_hep_ct = make_ct(cond_belum & (df_clean[c_dx] == "VIRAL HEPATITIS"))

  cond_dn = (df_clean[c_me] == target_me) & (
      df_clean[c_br].astype(str).str.strip().str.upper() == "DAFTAR NOTIFIKASI"
  )
  df_dn_ct = make_ct(cond_dn & df_clean[c_dx].isin(inclusion_tuple))
  df_dn_exc_ct = make_ct(cond_dn & ~df_clean[c_dx].isin(exclusion_tuple))
  df_hep_dn_ct = make_ct(cond_dn & (df_clean[c_dx] == "VIRAL HEPATITIS"))

  df_lewat = df_clean[df_clean[c_me] == target_me].copy()
  df_lewat = df_lewat[
      ~df_lewat[c_bp]
      .fillna("")
      .astype(str)
      .str.upper()
      .str.contains("PKD|PEJABAT")
  ].copy()

  diff = []
  for bc, bi in zip(
      df_lewat[c_bc].apply(parse_excel_date),
      df_lewat[c_bi].apply(parse_excel_date),
  ):
    diff.append((bc - bi).days if bc and bi else None)
  df_lewat["DIFF_DAYS"] = diff
  df_lewat["IS_SWASTA"] = (
      df_lewat[c_cd].fillna("").astype(str).str.upper().str.contains("SWASTA")
  )

  def agg_lewat(d_sub):
    g = {}
    for _, r in d_sub.iterrows():
      dx, dt, sw = r[c_dx], r[c_dt], r["IS_SWASTA"]
      if dx not in g:
        g[dx] = {d: {"t": 0, "s": 0} for d in VALID_DISTRICTS}
      if dt in g[dx]:
        g[dx][dt]["t"] += 1
        if sw:
          g[dx][dt]["s"] += 1
    res = []
    for dx, dists in g.items():
      tot = sum(dists[d]["t"] for d in VALID_DISTRICTS)
      if tot > 0:
        row = {
            "DIAGNOSIS": dx,
            "TOTAL": tot,
            "SWASTA": sum(dists[d]["s"] for d in VALID_DISTRICTS),
        }
        for d in VALID_DISTRICTS:
          row.update({f"{d}_tot": dists[d]["t"], f"{d}_swasta": dists[d]["s"]})
        res.append(row)
    return (
        pd.DataFrame(res)
        .sort_values(by=["TOTAL", "DIAGNOSIS"], ascending=[False, True])
        .reset_index(drop=True)
        if res
        else pd.DataFrame()
    )

  df_24h_raw = df_lewat[
      df_lewat[c_dx].isin(DIAG_TEMPOH_24H) & (df_lewat["DIFF_DAYS"] > 2)
  ]
  df_7d_raw = df_lewat[
      df_lewat[c_dx].isin(DIAG_TEMPOH_7D) & (df_lewat["DIFF_DAYS"] > 7)
  ]

  return (
      len(df),
      stats_semasa,
      stats_kumulatif,
      df_penyakit,
      df_district,
      df_belum_ct,
      df_hep_ct,
      df_dn_ct,
      df_dn_exc_ct,
      df_hep_dn_ct,
      agg_lewat(df_24h_raw),
      agg_lewat(df_7d_raw),
      df_24h_raw,
      df_7d_raw,
      fetch_wabak_data(),
      fetch_bencana_data(),
  )


def generate_lewat_excel(df_24h, df_7d):
  out = io.BytesIO()
  with pd.ExcelWriter(out, engine="openpyxl") as w:
    df_24h.to_excel(w, sheet_name="Lewat 24 Jam", index=False)
    df_7d.to_excel(w, sheet_name="Lewat 7 Hari", index=False)
  out.seek(0)
  return out


# ---------------------------------------------------------
# Presentation Generator
# ---------------------------------------------------------
def generate_pptx(
    stats_semasa,
    stats_kumulatif,
    df_penyakit,
    df_district,
    df_belum_ct,
    df_hep_ct,
    df_dn_ct,
    df_dn_exc_ct,
    df_hep_dn_ct,
    lewat_24h_df,
    lewat_7d_df,
    df_wabak,
    df_bencana,
    df_graf,
    df_ili_sari,
    df_sampel,
    df_prestasi,
    epi_week,
    year,
):
  try:
    prs = Presentation(fetch_google_slides_pptx(GOOGLE_SLIDES_ID))
  except:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)

  for slide in prs.slides:
    for shape in slide.shapes:
      if shape.has_table:
        for row in shape.table.rows:
          for cell in row.cells:
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            for p in cell.text_frame.paragraphs:
              p.alignment = PP_ALIGN.CENTER
              for run in p.runs:
                run.font.size = Pt(11)

  def add_bottom_banner(sl):
    b = sl.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(6.5),
        Inches(6.9),
        Inches(6.833),
        Inches(0.4),
    )
    b.fill.solid()
    b.fill.fore_color.rgb = NAVY
    b.line.fill.background()
    p = b.text_frame.paragraphs[0]
    p.text = (
        "UNIT SURVELAN & KESIAPSIAGAAN, JABATAN KESIHATAN NEGERI SELANGOR"
    )
    (
        p.font.size,
        p.font.color.rgb,
        p.alignment,
        p.font.name,
        p.font.bold,
    ) = (
        Pt(9),
        RGBColor(255, 255, 255),
        PP_ALIGN.RIGHT,
        "Calibri",
        True,
    )
    b.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE

  def add_slide_header(sl, t1, t2):
    script_dir = os.path.dirname(os.path.abspath(__file__))
    logo_path = (
        "logo.png"
        if os.path.exists("logo.png")
        else os.path.join(script_dir, "logo.png")
    )
    if os.path.exists(logo_path):
      sl.shapes.add_picture(logo_path, Inches(0.6), Inches(0.3), width=Inches(1.5))

    l = sl.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(2.3), Inches(0.4), Inches(0.03), Inches(0.9)
    )
    l.fill.solid()
    l.fill.fore_color.rgb = NAVY
    l.line.fill.background()
    tb = sl.shapes.add_textbox(
        Inches(2.4), Inches(0.35), Inches(10), Inches(0.9)
    )
    p = tb.text_frame.paragraphs[0]
    p.text = t1

    if "Tanpa Wabak Vektor" in t1:
      p.font.size = Pt(20)
    elif (
        "Sampel Survelan" in t1
        or "Sentinel Selangor" in t1
        or "Tren Keputusan" in t1
    ):
      p.font.size = Pt(19)
    elif (
        "Kadar konsultasi ILI & SARI" in t1
        or "Kluster Influenza" in t1
        or "Prestasi Penghantaran" in t1
    ):
      p.font.size = Pt(24)
    elif any(x in t1 for x in ["Senarai", "Bilangan", "Tren"]):
      p.font.size = Pt(28)
    else:
      p.font.size = Pt(36)

    p.font.bold = True
    p.font.color.rgb = NAVY

    p2 = tb.text_frame.add_paragraph()
    p2.text = t2
    p2.font.size = Pt(14) if "Tempoh" in t2 else Pt(16)
    p2.font.bold = True
    p2.font.color.rgb = NAVY

  # --- Slide 1: Title Slide ---
  slide1 = (
      prs.slides[0]
      if len(prs.slides) > 0
      else prs.slides.add_slide(prs.slide_layouts[6])
  )
  if len(prs.slides) > 0:
    for shape in list(slide1.shapes):
      shape._element.getparent().remove(shape._element)

  for x, y, w, h in [
      (0.2, 0.2, 2.2, 0.25),
      (0.2, 0.2, 0.25, 1.8),
      (10.933, 0.2, 2.2, 0.25),
      (12.883, 0.2, 0.25, 1.8),
      (0.2, 7.05, 2.2, 0.25),
      (0.2, 5.5, 0.25, 1.8),
      (10.933, 7.05, 2.2, 0.25),
      (12.883, 5.5, 0.25, 1.8),
  ]:
    s = slide1.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    s.fill.solid()
    s.fill.fore_color.rgb = NAVY
    s.line.fill.background()

  card = slide1.shapes.add_shape(
      MSO_SHAPE.RECTANGLE,
      Inches(0.42),
      Inches(0.42),
      Inches(12.493),
      Inches(6.66),
  )
  card.fill.solid()
  card.fill.fore_color.rgb = RGBColor(255, 255, 255)
  card.line.color.rgb = RGBColor(210, 210, 210)
  card.line.width = Pt(1.5)

  script_dir = os.path.dirname(os.path.abspath(__file__))
  logo_file = (
      "logo.png"
      if os.path.exists("logo.png")
      else os.path.join(script_dir, "logo.png")
  )
  if os.path.exists(logo_file):
    slide1.shapes.add_picture(
        logo_file, Inches(5.66), Inches(0.85), width=Inches(2.0)
    )

  tb = slide1.shapes.add_textbox(
      Inches(1.5), Inches(2.75), Inches(10.333), Inches(1.1)
  )
  p = tb.text_frame.paragraphs[0]
  p.text = "Selangor Epidemiology Review"
  (
      p.font.size,
      p.font.bold,
      p.font.name,
      p.font.color.rgb,
      p.alignment,
  ) = (
      Pt(50),
      True,
      "Calibri",
      NAVY,
      PP_ALIGN.CENTER,
  )

  l = slide1.shapes.add_shape(
      MSO_SHAPE.RECTANGLE, Inches(5.916), Inches(3.95), Inches(1.5), Inches(0.03)
  )
  l.fill.solid()
  l.fill.fore_color.rgb = RGBColor(200, 200, 200)
  l.line.fill.background()

  tb2 = slide1.shapes.add_textbox(
      Inches(1.5), Inches(4.1), Inches(10.333), Inches(0.7)
  )
  p2 = tb2.text_frame.paragraphs[0]
  p2.text = f"ME {epi_week:02d} / {year}"
  (
      p2.font.size,
      p2.font.bold,
      p2.font.name,
      p2.font.color.rgb,
      p2.alignment,
  ) = (
      Pt(28),
      True,
      "Calibri",
      NAVY,
      PP_ALIGN.CENTER,
  )

  # --- Add QR Code onto Title Slide ---
  qr_stream = None
  qr_file = (
      "qr.png"
      if os.path.exists("qr.png")
      else os.path.join(script_dir, "qr.png")
  )

  if os.path.exists(qr_file):
    qr_stream = qr_file
  else:
    try:
      qr_bytes = base64.b64decode(EMBEDDED_QR_B64)
      qr_stream = io.BytesIO(qr_bytes)
    except Exception:
      qr_stream = None

  if qr_stream:
    qr_size = 1.6
    qr_x = (13.333 - qr_size) / 2
    qr_y = 4.95
    slide1.shapes.add_picture(
        qr_stream,
        Inches(qr_x),
        Inches(qr_y),
        width=Inches(qr_size),
        height=Inches(qr_size),
    )

  # --- Slide 2: Analisa e-Notifikasi ---
  s2 = prs.slides.add_slide(prs.slide_layouts[6])
  add_slide_header(s2, "Analisa e-Notifikasi", f"ME {epi_week:02d} /{year}")
  tb = s2.shapes.add_table(
      7, 5, Inches(1.15), Inches(1.7), Inches(11.0), Inches(4.7)
  ).table
  for c in range(5):
    tb.columns[c].width = Inches(2.2)
  for i, txt in enumerate([
      "Minggu Epid",
      f"ME {epi_week:02d} (Semasa)",
      "",
      f"Kumulatif Sehingga ME {epi_week:02d}",
      "",
  ]):
    write_cell(tb.cell(0, i), txt, size=17)
  tb.cell(0, 1).merge(tb.cell(0, 2))
  tb.cell(0, 3).merge(tb.cell(0, 4))
  for i, label_txt in enumerate([
      "",
      "Jumlah Notifikasi",
      "Daftar Notifikasi",
      "Daftar Kes",
      "Abai Notifikasi",
      "Belum Ambil Tindakan",
      "Batal Daftar",
  ]):
    if i > 0:
      write_cell(tb.cell(i, 0), label_txt, size=17)

  write_cell(tb.cell(1, 1), f"{stats_semasa['total']:,}", size=17)
  write_cell(tb.cell(1, 3), f"{stats_kumulatif['total']:,}", size=17)

  pct_key_map = {
      "daftar_notifikasi": "pct_daftar_notif",
      "daftar_kes": "pct_daftar_kes",
      "abai": "pct_abai",
      "belum": "pct_belum",
      "batal": "pct_batal",
  }

  for r, k in enumerate(
      ["daftar_notifikasi", "daftar_kes", "abai", "belum", "batal"], 2
  ):
    write_cell(tb.cell(r, 1), f"{stats_semasa[k]:,}", size=17)
    write_cell(tb.cell(r, 2), stats_semasa[pct_key_map[k]], size=17)
    write_cell(tb.cell(r, 3), f"{stats_kumulatif[k]:,}", size=17)
    write_cell(tb.cell(r, 4), stats_kumulatif[pct_key_map[k]], size=17)

  for i, row in enumerate(tb.rows):
    for j, cell in enumerate(row.cells):
      set_cell_border(cell)
      cell.vertical_anchor = MSO_ANCHOR.MIDDLE
      if i == 0 or j == 0:
        cell.fill.solid()
        cell.fill.fore_color.rgb = LIGHT_GREY
      else:
        cell.fill.solid()
        cell.fill.fore_color.rgb = (
            RGBColor(0, 0, 0)
            if i == 1 and j in [2, 4]
            else RGBColor(255, 255, 255)
        )

  timestamp = (
      datetime.datetime.now(datetime.timezone(datetime.timedelta(hours=8)))
      .strftime("%d/%m/%Y @ %I.%M%p")
      .upper()
  )
  t_box = s2.shapes.add_textbox(
      Inches(0.1), Inches(6.9), Inches(10), Inches(0.4)
  ).text_frame.paragraphs[0]
  t_box.text = (
      f"(Sumber : Sistem e-notifikasi, KKM muat turun pada ({timestamp}))"
  )
  t_box.font.size, t_box.font.italic, t_box.font.bold = Pt(9), True, True
  add_bottom_banner(s2)

  # --- Slide 3: Bilangan Daftar Kes ---
  s3 = prs.slides.add_slide(prs.slide_layouts[6])
  add_slide_header(
      s3, "Bilangan Daftar Kes Mengikut Penyakit", f"ME {epi_week:02d} / {year}"
  )
  df_r = (
      df_penyakit.head(19)
      if not df_penyakit.empty and len(df_penyakit) > 19
      else df_penyakit
  )
  tb3 = s3.shapes.add_table(
      len(df_r) + 2, 4, Inches(0.8), Inches(1.5), Inches(11.733), Inches(5.0)
  ).table
  (
      tb3.columns[0].width,
      tb3.columns[1].width,
      tb3.columns[2].width,
      tb3.columns[3].width,
  ) = (Inches(3.2), Inches(2.2), Inches(2.7), Inches(3.633))

  for j, h in enumerate([
      "Penyakit",
      f"ME {epi_week:02d}",
      f"≤ ME {epi_week:02d}/ {year}",
      "Peratus Daftar Kes\n(Bil Daftar Kes / Bil Notifikasi Kes Tersebut)",
  ]):
    write_cell(tb3.cell(0, j), h)
  for i, (_, r) in enumerate(df_r.iterrows()):
    r_i = i + 1
    write_cell(tb3.cell(r_i, 0), str(r["Penyakit"]), align_left=True)
    write_cell(tb3.cell(r_i, 1), f"{int(r['Semasa']):,}")
    p = write_cell(tb3.cell(r_i, 2), f"{int(r['Kumulatif']):,}")
    if r["Mati"] > 0:
      run = p.add_run()
      run.text = f" ({int(r['Mati'])})"
      run.font.color.rgb = RGBColor(255, 0, 0)
      run.font.size = Pt(10)
      run.font.bold = True
    write_cell(tb3.cell(r_i, 3), f"{int(r['Peratus'])}%")

  r_i = len(tb3.rows) - 1
  write_cell(tb3.cell(r_i, 0), "JUMLAH")
  write_cell(
      tb3.cell(r_i, 1),
      f"{int(df_penyakit['Semasa'].sum() if not df_penyakit.empty else 0):,}",
  )
  p = write_cell(
      tb3.cell(r_i, 2),
      f"{int(df_penyakit['Kumulatif'].sum() if not df_penyakit.empty else 0):,}",
  )
  if (
      m := df_penyakit["Mati"].sum() if not df_penyakit.empty else 0
  ) > 0:
    run = p.add_run()
    run.text = f" ({int(m)})"
    run.font.color.rgb = RGBColor(255, 0, 0)
    run.font.size = Pt(10)
    run.font.bold = True

  for i, row in enumerate(tb3.rows):
    for j, cell in enumerate(row.cells):
      set_cell_border(cell)
      cell.vertical_anchor = MSO_ANCHOR.MIDDLE
      cell.fill.solid()
      cell.fill.fore_color.rgb = (
          LIGHT_GREY
          if i == 0 or i == len(tb3.rows) - 1
          else RGBColor(255, 255, 255)
      )

  tf3 = s3.shapes.add_textbox(
      Inches(0.8), Inches(6.9), Inches(6), Inches(0.4)
  ).text_frame
  tf3.paragraphs[0].text = "Sumber data adalah daripada sistem eNotifikasi"
  tf3.paragraphs[0].font.size, tf3.paragraphs[0].font.bold = Pt(9), True
  p2 = tf3.add_paragraph()
  p2.text = "*(Mati)"
  p2.font.size, p2.font.bold, p2.font.color.rgb = (
      Pt(9),
      True,
      RGBColor(255, 0, 0),
  )
  add_bottom_banner(s3)

  # --- Slide 4: Daerah ---
  s4 = prs.slides.add_slide(prs.slide_layouts[6])
  add_slide_header(
      s4, "Status pencapaian e-Notifikasi", f"Sehingga ME {epi_week:02d} / {year}"
  )
  tb4 = s4.shapes.add_table(
      len(df_district) + 2,
      12,
      Inches(0.5),
      Inches(1.5),
      Inches(12.333),
      Inches(5.0),
  ).table
  for j, w in enumerate(
      [1.4, 1.1, 1.0, 0.85, 0.95, 0.85, 1.0, 0.85, 0.95, 0.85, 1.0, 0.85]
  ):
    tb4.columns[j].width = Inches(w)

  for j, h in enumerate([
      "DAERAH",
      "Jumlah\nNotifikasi",
      "Daftar\nNotifikasi",
      "%",
      "Daftar Kes",
      "%",
      "Abai\nNotifikasi",
      "%",
      "Batal Daftar",
      "%",
      "Belum\nAmbil\nTindakan",
      "%",
  ]):
    write_cell(tb4.cell(0, j), h)

  for i, (_, r) in enumerate(df_district.iterrows()):
    ri = i + 1
    write_cell(tb4.cell(ri, 0), str(r["DAERAH"]))
    for j, k in enumerate(
        [
            "Jumlah Notifikasi",
            "Daftar Notifikasi",
            "pct_notif",
            "Daftar Kes",
            "pct_kes",
            "Abai Notifikasi",
            "pct_abai",
            "Batal Daftar",
            "pct_batal",
            "Belum Ambil Tindakan",
            "pct_belum",
        ],
        1,
    ):
      write_cell(
          tb4.cell(ri, j),
          f"{r[k]:.2f}%" if "pct" in k else f"{int(r[k]):,}",
      )

  ri = len(tb4.rows) - 1
  tj = df_district["Jumlah Notifikasi"].sum() if not df_district.empty else 0
  write_cell(tb4.cell(ri, 0), "JUMLAH")
  write_cell(tb4.cell(ri, 1), f"{int(tj):,}")
  for j, k in zip(
      range(2, 12, 2),
      [
          "Daftar Notifikasi",
          "Daftar Kes",
          "Abai Notifikasi",
          "Batal Daftar",
          "Belum Ambil Tindakan",
      ],
  ):
    tv = df_district[k].sum() if not df_district.empty else 0
    write_cell(tb4.cell(ri, j), f"{int(tv):,}")
    write_cell(tb4.cell(ri, j + 1), f"{(tv/tj*100):.2f}%" if tj else "0.00%")

  for i, row in enumerate(tb4.rows):
    for j, cell in enumerate(row.cells):
      set_cell_border(cell)
      cell.vertical_anchor = MSO_ANCHOR.MIDDLE
      cell.fill.solid()
      cell.fill.fore_color.rgb = (
          LIGHT_GREY
          if i == 0 or j == 0 or i == len(tb4.rows) - 1
          else RGBColor(255, 255, 255)
      )
  add_bottom_banner(s4)

  # --- Slide 5 to 9: CT Tables ---
  def build_ct_slide(title, subtitle, df_ct):
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    add_slide_header(sl, title, subtitle)
    tb = sl.shapes.add_table(
        len(df_ct) + 3,
        11,
        Inches(0.5),
        Inches(1.5),
        Inches(12.333),
        Inches(5.0),
    ).table
    tb.columns[0].width = Inches(2.333)
    for j in range(1, 11):
      tb.columns[j].width = Inches(1.0)

    for j, txt in enumerate(
        ["DIAGNOSIS"] + [d[1] for d in DISTRICT_ABBR] + ["JUM"]
    ):
      write_cell(tb.cell(0, j), txt, size=16)
    for j in range(1, 11):
      write_cell(tb.cell(1, j), f"ME{epi_week:02d}", size=16)
    tb.cell(0, 0).merge(tb.cell(1, 0))
    write_cell(tb.cell(0, 0), "DIAGNOSIS", size=16)

    ds = {d[0]: 0 for d in DISTRICT_ABBR}
    tot = 0
    cn = df_ct.columns[0] if not df_ct.empty else "DIAGNOSIS"

    for i, (_, r) in enumerate(df_ct.iterrows()):
      ri = i + 2
      rj = 0
      write_cell(tb.cell(ri, 0), str(r[cn]), align_left=True, size=16)
      for j, d_name in enumerate([d[0] for d in DISTRICT_ABBR]):
        v = int(r[d_name]) if d_name in r else 0
        write_cell(tb.cell(ri, j + 1), str(v), size=16)
        ds[d_name] += v
        rj += v
      write_cell(
          tb.cell(ri, 10), str(int(r["JUM"]) if "JUM" in r else rj), size=16
      )
      tot += int(r["JUM"]) if "JUM" in r else rj

    ri = len(tb.rows) - 1
    write_cell(tb.cell(ri, 0), "JUMLAH", size=16)
    for j, d_name in enumerate([d[0] for d in DISTRICT_ABBR]):
      write_cell(tb.cell(ri, j + 1), str(ds[d_name]), size=16)
    write_cell(tb.cell(ri, 10), str(tot), size=16)

    for i, row in enumerate(tb.rows):
      for j, cell in enumerate(row.cells):
        set_cell_border(cell)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.fill.solid()
        cell.fill.fore_color.rgb = (
            LIGHT_GREY
            if i in [0, 1] or j == 0 or i == len(tb.rows) - 1
            else RGBColor(255, 255, 255)
        )
    add_bottom_banner(sl)

  build_ct_slide(
      "Senarai Kes Belum Ambil Tindakan Mengikut Diagnosis",
      f"ME {epi_week:02d} / {year}",
      df_belum_ct,
  )
  build_ct_slide(
      "Senarai Kes Belum Ambil Tindakan Mengikut Diagnosis",
      f"ME {epi_week:02d} / {year}",
      df_hep_ct,
  )
  build_ct_slide(
      "Bilangan Kes Yang Masih Berstatus Daftar Notifikasi",
      "Tempoh daftar kes ≤ 7 hari",
      df_dn_ct,
  )
  build_ct_slide(
      "Bilangan Kes Yang Masih Berstatus Daftar Notifikasi",
      "Tempoh daftar kes ≤ 14 hari",
      df_dn_exc_ct,
  )
  build_ct_slide(
      "Bilangan Kes Yang Masih Berstatus Daftar Notifikasi",
      "Tempoh daftar kes ≤ 14 hari",
      df_hep_dn_ct,
  )

  # --- Slide 10 & 11: Lewat Notifikasi ---
  def build_lewat_slide(title, df_l):
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    add_slide_header(sl, title, f"ME {epi_week:02d} / {year}")
    tb = sl.shapes.add_table(
        len(df_l) + 2, 11, Inches(0.5), Inches(1.5), Inches(12.333), Inches(4.5)
    ).table
    tb.columns[0].width = Inches(2.0)
    for j in range(1, 11):
      tb.columns[j].width = Inches(1.033)
    for j, h in enumerate([
        "Diagnosis",
        "GOMBAK",
        "HULU\nLANGAT",
        "HULU\nSELANGOR",
        "KLANG",
        "KUALA\nLANGAT",
        "KUALA\nSELANGOR",
        "PETALING",
        "SABAK\nBERNAM",
        "SEPANG",
        "JUMLAH",
    ]):
      write_cell(tb.cell(0, j), h, size=12)

    ds = {d: {"t": 0, "s": 0} for d in VALID_DISTRICTS}
    tt = 0
    ts = 0
    for i, (_, r) in enumerate(df_l.iterrows()):
      ri = i + 1
      write_cell(tb.cell(ri, 0), str(r["DIAGNOSIS"]), align_left=True, size=12)
      for j, d in enumerate(VALID_DISTRICTS):
        t, s = int(r[f"{d}_tot"]), int(r[f"{d}_swasta"])
        ds[d]["t"] += t
        ds[d]["s"] += s
        write_cell(tb.cell(ri, j + 1), format_cell_stat(t, s, True), size=12)
      tt += int(r["TOTAL"])
      ts += int(r["SWASTA"])
      write_cell(
          tb.cell(ri, 10),
          format_cell_stat(int(r["TOTAL"]), int(r["SWASTA"]), True),
          size=12,
      )

    ri = len(tb.rows) - 1
    write_cell(tb.cell(ri, 0), "JUMLAH", size=12)
    for j, d in enumerate(VALID_DISTRICTS):
      write_cell(
          tb.cell(ri, j + 1),
          format_cell_stat(ds[d]["t"], ds[d]["s"], False),
          size=12,
      )
    write_cell(tb.cell(ri, 10), format_cell_stat(tt, ts, False), size=12)

    for i, row in enumerate(tb.rows):
      for j, cell in enumerate(row.cells):
        set_cell_border(cell)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.fill.solid()
        cell.fill.fore_color.rgb = (
            LIGHT_GREY
            if i == 0 or j == 0 or i == len(tb.rows) - 1
            else RGBColor(255, 255, 255)
        )

    tf = sl.shapes.add_textbox(
        Inches(0.5), Inches(6.85), Inches(6), Inches(0.4)
    ).text_frame
    tf.paragraphs[0].text = "( ) Kemudahan Kesihatan Swasta"
    tf.paragraphs[0].font.size, tf.paragraphs[0].font.bold = Pt(11), True
    add_bottom_banner(sl)

  build_lewat_slide("Lewat Notifikasi 24 Jam", lewat_24h_df)
  build_lewat_slide("Lewat Notifikasi 7 Hari", lewat_7d_df)

  # --- Slide 12+: Wabak Mengikut Daerah ---
  if not df_wabak.empty:
    for p_i in range((len(df_wabak) + 9) // 10):
      chunk = df_wabak.iloc[p_i * 10 : (p_i + 1) * 10]
      sl = prs.slides.add_slide(prs.slide_layouts[6])
      add_slide_header(
          sl,
          "Bilangan Keseluruhan Episod Wabak / Kluster Mengikut Daerah",
          f"Sehingga ME {epi_week:02d} / {year}",
      )

      tb = sl.shapes.add_table(
          len(chunk) + 2,
          12,
          Inches(0.4),
          Inches(1.5),
          Inches(12.533),
          Inches(5.0),
      ).table
      tb.columns[0].width, tb.columns[1].width, tb.columns[11].width = (
          Inches(0.7),
          Inches(2.4),
          Inches(1.333),
      )
      for c in range(2, 11):
        tb.columns[c].width = Inches(0.9)

      tb.cell(0, 2).merge(tb.cell(0, 10))
      tb.cell(0, 0).merge(tb.cell(1, 0))
      tb.cell(0, 1).merge(tb.cell(1, 1))
      tb.cell(0, 11).merge(tb.cell(1, 11))
      write_wabak_cell(tb.cell(0, 0), "Bil")
      write_wabak_cell(tb.cell(0, 1), "Wabak")
      write_header_with_red_me(
          tb.cell(0, 2),
          "Pecahan kumulatif mengikut daerah",
          f"(ME {epi_week:02d} / {year})",
      )
      write_header_with_red_me(
          tb.cell(0, 11),
          "Kumulatif sehingga ME",
          f"({epi_week:02d} / {year})",
          newline=True,
      )
      for ci, ab in enumerate(
          ["GBK", "HL", "HS", "KLG", "KL", "KS", "PTG", "SB", "SPG"]
      ):
        write_wabak_cell(tb.cell(1, ci + 2), ab)

      for ri, (_, r) in enumerate(chunk.iterrows()):
        cr = ri + 2
        wn = str(r.iloc[1]).strip()
        ij = wn.upper() == "JUMLAH"
        write_wabak_cell(tb.cell(cr, 0), "" if ij else str(p_i * 10 + ri + 1))
        write_wabak_cell(tb.cell(cr, 1), wn, align_left=True)
        for ci in range(9):
          write_wabak_cell(
              tb.cell(cr, ci + 2), r.iloc[ci + 2] if (ci + 2) < len(r) else "-"
          )
        write_wabak_cell(tb.cell(cr, 11), r.iloc[11] if 11 < len(r) else "-")

      for i, row in enumerate(tb.rows):
        ij = i >= 2 and row.cells[1].text_frame.text.strip().upper() == "JUMLAH"
        for j, cell in enumerate(row.cells):
          set_cell_border(cell)
          cell.vertical_anchor = MSO_ANCHOR.MIDDLE
          cell.fill.solid()
          cell.fill.fore_color.rgb = (
              LIGHT_GREY
              if i in [0, 1] or j in [0, 1] or ij
              else RGBColor(255, 255, 255)
          )
      add_bottom_banner(sl)

  # --- Slide 13+: Insiden/Bencana ---
  if not df_bencana.empty:
    for p_i in range((len(df_bencana) + 7) // 8):
      chunk = df_bencana.iloc[p_i * 8 : (p_i + 1) * 8]
      sl = prs.slides.add_slide(prs.slide_layouts[6])
      add_slide_header(
          sl,
          "Bilangan Insiden/Bencana Mengikut Daerah",
          f"Sehingga ME {epi_week:02d} / {year}",
      )

      tb = sl.shapes.add_table(
          len(chunk) + 2,
          14,
          Inches(0.4),
          Inches(1.5),
          Inches(12.533),
          Inches(4.5),
      ).table
      tb.columns[0].width = Inches(1.8)
      for c in range(1, 14):
        tb.columns[c].width = Inches(0.82)

      tb.cell(0, 1).merge(tb.cell(0, 11))
      tb.cell(0, 0).merge(tb.cell(1, 0))
      tb.cell(0, 12).merge(tb.cell(1, 12))
      tb.cell(0, 13).merge(tb.cell(1, 13))
      write_wabak_cell(tb.cell(0, 0), "INSIDEN/BENCANA", size=13)
      write_header_with_red_me(
          tb.cell(0, 1),
          "Pecahan kumulatif mengikut daerah",
          f"(ME {epi_week:02d} / {year})",
          size=13,
      )
      write_wabak_cell(tb.cell(0, 12), "JUMLAH", size=13)
      write_wabak_cell(tb.cell(0, 13), "DIISYTIHAR OLEH CPRC KKM", size=11)
      for ci, ab in enumerate([
          "GOMBAK",
          "HULU\nLANGAT",
          "HULU\nSELANGOR",
          "KLANG",
          "KUALA\nLANGAT",
          "KUALA\nSELANGOR",
          "PETALING",
          "SABAK\nBERNAM",
          "SEPANG",
          "PK P.KLANG",
          "PK KLIA",
      ]):
        write_wabak_cell(tb.cell(1, ci + 1), ab, size=11)

      for ri, (_, r) in enumerate(chunk.iterrows()):
        cr = ri + 2
        in_name = str(r.iloc[0]).strip()
        ij = in_name.upper() == "JUMLAH"
        write_wabak_cell(tb.cell(cr, 0), in_name, align_left=True, size=13)
        for ci in range(1, 14):
          write_wabak_cell(
              tb.cell(cr, ci), r.iloc[ci] if ci < len(r) else "-", size=13
          )

      for i, row in enumerate(tb.rows):
        ij = i >= 2 and row.cells[0].text_frame.text.strip().upper() == "JUMLAH"
        for j, cell in enumerate(row.cells):
          set_cell_border(cell)
          cell.vertical_anchor = MSO_ANCHOR.MIDDLE
          cell.fill.solid()
          cell.fill.fore_color.rgb = (
              LIGHT_GREY
              if i in [0, 1] or j == 0 or ij
              else RGBColor(255, 255, 255)
          )
      add_bottom_banner(sl)

  # Common Trend Chart Data Truncation Logic
  df_graf_clean = pd.DataFrame()
  if not df_graf.empty and "Minggu Epid" in df_graf.columns:
    has_reset = False
    keep_indices = []
    for idx, row in df_graf.iterrows():
      try:
        w_num = int(row["Minggu Epid"])
        if idx > 0:
          prev_w_num = int(df_graf.iloc[idx - 1]["Minggu Epid"])
          if w_num < prev_w_num and prev_w_num >= 40:
            has_reset = True

        if has_reset and w_num > epi_week:
          break
        keep_indices.append(idx)
      except ValueError:
        keep_indices.append(idx)

    df_graf_clean = df_graf.loc[keep_indices].reset_index(drop=True)

  chart_subtitle = f"ME01 /{year-1} - ME {epi_week:02d} /{year}"

  # --- Slide X: Tren Wabak Native Chart (All Diseases) ---
  if not df_graf_clean.empty:
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    add_slide_header(
        sl, "Tren Wabak Mengikut Jenis Penyakit Berjangkit", chart_subtitle
    )

    categories = [str(x) for x in df_graf_clean["Minggu Epid"].tolist()]
    series_cols = [
        c
        for c in df_graf_clean.columns
        if c != "Minggu Epid"
        and str(c).lower().strip() not in ["nan", "none", "", "null"]
        and not str(c).startswith("Unnamed")
    ]

    chart_data = CategoryChartData()
    chart_data.categories = categories
    for sn in series_cols:
      series_vals = (
          pd.to_numeric(df_graf_clean[sn], errors="coerce").fillna(0).tolist()
      )
      chart_data.add_series(str(sn), series_vals)

    chart = sl.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_STACKED_100,
        Inches(0.4),
        Inches(1.5),
        Inches(9.2),
        Inches(5.1),
        chart_data,
    ).chart

    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.RIGHT
    chart.legend.include_in_layout = False
    chart.legend.font.size = Pt(8.5)
    chart.legend.font.name = "Calibri"
    chart.legend.font.bold = True

    val_axis = chart.value_axis
    val_axis.has_major_gridlines = False
    val_axis.has_minor_gridlines = False
    val_axis.format.line.fill.background()
    val_axis.tick_labels.font.size = Pt(10)
    val_axis.tick_labels.font.name = "Calibri"
    val_axis.tick_labels.font.bold = True

    cat_axis = chart.category_axis
    cat_axis.format.line.fill.background()
    cat_axis.tick_labels.font.size = Pt(10)
    cat_axis.tick_labels.font.name = "Calibri"
    cat_axis.tick_labels.font.bold = True

    plotArea = chart.element.find(qn("c:chart")).find(qn("c:plotArea"))
    for ax_tag in [qn("c:catAx"), qn("c:valAx")]:
      for ax_elem in plotArea.findall(ax_tag):
        remove_axis_line(ax_elem)

    try:
      chart.plots[0].gap_width = 20
    except:
      pass

    for series in chart.series:
      sn = series.name.strip()
      color = DISEASE_COLORS.get(sn, None)
      if not color:
        for key, c in DISEASE_COLORS.items():
          if key.lower() in sn.lower() or sn.lower() in key.lower():
            color = c
            break
      if color:
        series.format.fill.solid()
        series.format.fill.fore_color.rgb = color

    idx_reset = -1
    for i in range(1, len(categories)):
      try:
        c_curr = int(categories[i])
        c_prev = int(categories[i - 1])
        if c_curr < c_prev and c_prev >= 40:
          idx_reset = i
          break
      except ValueError:
        pass

    if idx_reset != -1:
      plot_width = 9.2
      x_pos = 0.4 + (plot_width / len(categories)) * idx_reset
      line = sl.shapes.add_connector(
          MSO_CONNECTOR.STRAIGHT,
          Inches(x_pos),
          Inches(1.5),
          Inches(x_pos),
          Inches(6.5),
      )
      line.line.color.rgb = RGBColor(0, 0, 0)
      line.line.width = Pt(4)
      line.line.dash_style = 7

    add_bottom_banner(sl)

  # --- Slide X+1: Tren Wabak Native Chart (Tanpa Wabak Vektor) ---
  if not df_graf_clean.empty:
    sl_nv = prs.slides.add_slide(prs.slide_layouts[6])
    add_slide_header(
        sl_nv,
        "Tren Wabak Mengikut Jenis Penyakit Berjangkit (Tanpa Wabak Vektor)",
        chart_subtitle,
    )

    categories = [str(x) for x in df_graf_clean["Minggu Epid"].tolist()]

    vector_diseases_lower = ["denggi", "malaria", "chikungunya"]
    series_cols_non_vector = [
        c
        for c in df_graf_clean.columns
        if c != "Minggu Epid"
        and str(c).lower().strip() not in ["nan", "none", "", "null"]
        and not str(c).startswith("Unnamed")
        and str(c).lower().strip() not in vector_diseases_lower
    ]

    chart_data_nv = CategoryChartData()
    chart_data_nv.categories = categories
    for sn in series_cols_non_vector:
      series_vals = (
          pd.to_numeric(df_graf_clean[sn], errors="coerce").fillna(0).tolist()
      )
      chart_data_nv.add_series(str(sn), series_vals)

    chart_nv = sl_nv.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_STACKED_100,
        Inches(0.4),
        Inches(1.5),
        Inches(9.2),
        Inches(5.1),
        chart_data_nv,
    ).chart

    chart_nv.has_legend = True
    chart_nv.legend.position = XL_LEGEND_POSITION.RIGHT
    chart_nv.legend.include_in_layout = False
    chart_nv.legend.font.size = Pt(8.5)
    chart_nv.legend.font.name = "Calibri"
    chart_nv.legend.font.bold = True

    val_axis_nv = chart_nv.value_axis
    val_axis_nv.has_major_gridlines = False
    val_axis_nv.has_minor_gridlines = False
    val_axis_nv.format.line.fill.background()
    val_axis_nv.tick_labels.font.size = Pt(10)
    val_axis_nv.tick_labels.font.name = "Calibri"
    val_axis_nv.tick_labels.font.bold = True

    cat_axis_nv = chart_nv.category_axis
    cat_axis_nv.format.line.fill.background()
    cat_axis_nv.tick_labels.font.size = Pt(10)
    cat_axis_nv.tick_labels.font.name = "Calibri"
    cat_axis_nv.tick_labels.font.bold = True

    plotArea_nv = chart_nv.element.find(qn("c:chart")).find(qn("c:plotArea"))
    for ax_tag in [qn("c:catAx"), qn("c:valAx")]:
      for ax_elem in plotArea_nv.findall(ax_tag):
        remove_axis_line(ax_elem)

    try:
      chart_nv.plots[0].gap_width = 20
    except:
      pass

    for series in chart_nv.series:
      sn = series.name.strip()
      color = DISEASE_COLORS.get(sn, None)
      if not color:
        for key, c in DISEASE_COLORS.items():
          if key.lower() in sn.lower() or sn.lower() in key.lower():
            color = c
            break
      if color:
        series.format.fill.solid()
        series.format.fill.fore_color.rgb = color

    idx_reset = -1
    for i in range(1, len(categories)):
      try:
        c_curr = int(categories[i])
        c_prev = int(categories[i - 1])
        if c_curr < c_prev and c_prev >= 40:
          idx_reset = i
          break
      except ValueError:
        pass

    if idx_reset != -1:
      plot_width = 9.2
      x_pos = 0.4 + (plot_width / len(categories)) * idx_reset
      line_nv = sl_nv.shapes.add_connector(
          MSO_CONNECTOR.STRAIGHT,
          Inches(x_pos),
          Inches(1.5),
          Inches(x_pos),
          Inches(6.5),
      )
      line_nv.line.color.rgb = RGBColor(0, 0, 0)
      line_nv.line.width = Pt(4)
      line_nv.line.dash_style = 7

    add_bottom_banner(sl_nv)

  # --- Slide X+2: Tren Kluster Influenza dan Kadar konsultasi ILI & SARI ---
  df_ili_clean = pd.DataFrame()
  if not df_ili_sari.empty and "Minggu Epid" in df_ili_sari.columns:
    has_reset = False
    keep_indices = []
    for idx, row in df_ili_sari.iterrows():
      try:
        w_num = int(row["Minggu Epid"])
        if idx > 0:
          prev_w_num = int(df_ili_sari.iloc[idx - 1]["Minggu Epid"])
          if w_num < prev_w_num and prev_w_num >= 40:
            has_reset = True

        if has_reset and w_num > epi_week:
          break
        keep_indices.append(idx)
      except ValueError:
        keep_indices.append(idx)

    df_ili_clean = df_ili_sari.loc[keep_indices].reset_index(drop=True)

  if not df_ili_clean.empty:
    sl_is = prs.slides.add_slide(prs.slide_layouts[6])
    add_slide_header(
        sl_is,
        "Tren Kluster Influenza dan Kadar konsultasi ILI & SARI di Selangor",
        chart_subtitle,
    )

    categories = [str(x) for x in df_ili_clean["Minggu Epid"].tolist()]
    series_cols = [c for c in df_ili_clean.columns if c != "Minggu Epid"]

    chart_data_is = CategoryChartData()
    chart_data_is.categories = categories

    col_series_name = (
        series_cols[0] if len(series_cols) > 0 else "Bilangan Kluster ILI"
    )
    chart_data_is.add_series(
        str(col_series_name),
        (
            pd.to_numeric(df_ili_clean[col_series_name], errors="coerce")
            .fillna(0)
            .tolist()
        ),
    )

    for sn in series_cols[1:]:
      chart_data_is.add_series(
          str(sn),
          pd.to_numeric(df_ili_clean[sn], errors="coerce").fillna(0).tolist(),
      )

    chart_shape = sl_is.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        Inches(0.6),
        Inches(1.4),
        Inches(12.133),
        Inches(4.8),
        chart_data_is,
    )
    chart_is = chart_shape.chart

    # Configure Combo Chart:
    # Series 0 (Bilangan Kluster ILI - Bars) -> Primary Left Y-Axis (0 - 80)
    # Series 1 (Kadar Kemasukan SARI - Red Line) & Series 2 (Kadar Konsultasi ILI - Green Line) -> Secondary Right Y-Axis (0 - 30)
    def configure_combo_ili_sari(chart_obj):
      plotArea = chart_obj.element.find(qn("c:chart")).find(qn("c:plotArea"))
      barChart = plotArea.find(qn("c:barChart"))

      if barChart is None:
        return

      catAx_id = plotArea.find(qn("c:catAx")).find(qn("c:axId")).get("val")
      sec_valAx_id = "98765432"

      lineChart = OxmlElement("c:lineChart")
      lineChart.append(OxmlElement("c:grouping"))
      lineChart.find(qn("c:grouping")).set("val", "standard")

      for ser in list(barChart.findall(qn("c:ser"))):
        idx_val = int(ser.find(qn("c:idx")).get("val"))
        if idx_val in [1, 2]:
          barChart.remove(ser)

          marker = ser.find(qn("c:marker"))
          if marker is None:
            marker = OxmlElement("c:marker")
            ser.append(marker)
          symbol = marker.find(qn("c:symbol"))
          if symbol is None:
            symbol = OxmlElement("c:symbol")
            marker.append(symbol)
          symbol.set("val", "none")

          smooth = ser.find(qn("c:smooth"))
          if smooth is None:
            smooth = OxmlElement("c:smooth")
            ser.append(smooth)
          smooth.set("val", "0")

          lineChart.append(ser)

      axId1 = OxmlElement("c:axId")
      axId1.set("val", catAx_id)
      axId2 = OxmlElement("c:axId")
      axId2.set("val", sec_valAx_id)
      lineChart.append(axId1)
      lineChart.append(axId2)

      plotArea.insert(plotArea.index(barChart) + 1, lineChart)

      sec_valAx = OxmlElement("c:valAx")
      sec_axId = OxmlElement("c:axId")
      sec_axId.set("val", sec_valAx_id)
      sec_valAx.append(sec_axId)

      scaling = OxmlElement("c:scaling")
      scaling.append(OxmlElement("c:orientation"))
      scaling.find(qn("c:orientation")).set("val", "minMax")
      sec_valAx.append(scaling)

      delete = OxmlElement("c:delete")
      delete.set("val", "0")
      sec_valAx.append(delete)

      axPos = OxmlElement("c:axPos")
      axPos.set("val", "r")
      sec_valAx.append(axPos)

      tickLblPos = OxmlElement("c:tickLblPos")
      tickLblPos.set("val", "nextTo")
      sec_valAx.append(tickLblPos)

      crossAx = OxmlElement("c:crossAx")
      crossAx.set("val", catAx_id)
      sec_valAx.append(crossAx)

      crosses = OxmlElement("c:crosses")
      crosses.set("val", "autoZero")
      sec_valAx.append(crosses)

      plotArea.append(sec_valAx)

    if len(series_cols) > 1:
      configure_combo_ili_sari(chart_is)

    set_val_axis_title_oxml(
        chart_is, "Bilangan Kluster", axis_index=0, font_size=10, bold=True
    )

    val_axis_is = chart_is.value_axis
    val_axis_is.has_major_gridlines = False
    val_axis_is.has_minor_gridlines = False
    val_axis_is.format.line.fill.background()
    val_axis_is.tick_labels.font.size = Pt(10)
    val_axis_is.tick_labels.font.name = "Calibri"
    val_axis_is.tick_labels.font.bold = True

    cat_axis_is = chart_is.category_axis
    cat_axis_is.has_title = True
    cat_axis_is.axis_title.text_frame.text = "Minggu Epid"
    p_x_title = cat_axis_is.axis_title.text_frame.paragraphs[0]
    p_x_title.font.size = Pt(10)
    p_x_title.font.name = "Calibri"
    p_x_title.font.bold = True
    p_x_title.font.color.rgb = NAVY

    cat_axis_is.has_major_gridlines = False
    cat_axis_is.has_minor_gridlines = False
    cat_axis_is.format.line.fill.background()
    cat_axis_is.tick_labels.font.size = Pt(10)
    cat_axis_is.tick_labels.font.name = "Calibri"
    cat_axis_is.tick_labels.font.bold = True

    set_val_axis_title_oxml(
        chart_is,
        "Kadar Konsultasi ILI / Kemasukan Kes SARI",
        axis_index=1,
        font_size=10,
        bold=True,
    )

    chart_is.has_legend = True
    chart_is.legend.position = XL_LEGEND_POSITION.BOTTOM
    chart_is.legend.include_in_layout = False
    chart_is.legend.font.size =
